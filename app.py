import os
import re
import json
import base64
from pathlib import Path
from datetime import datetime
from flask import (Flask, render_template, request, redirect, url_for,
                   jsonify, send_from_directory, session, flash)
from markupsafe import escape as html_escape
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from flask_wtf.csrf import CSRFProtect

load_dotenv()

app = Flask(__name__)
csrf = CSRFProtect(app)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024  # 100 MB
_secret = os.environ.get("FLASK_SECRET_KEY")
if not _secret:
    if os.environ.get("VERCEL") or os.environ.get("FLASK_ENV") == "production":
        raise RuntimeError("FLASK_SECRET_KEY deve ser definida em produção!")
    import secrets
    _secret = secrets.token_hex(32)
app.secret_key = _secret

# Vercel tem filesystem read-only, usar /tmp
_TMP = Path("/tmp")
app.config["UPLOAD_FOLDER"] = _TMP / "uploads"
app.config["RESUMOS_FOLDER"] = _TMP / "resumos"

AUDIO_EXTS = {".mp3", ".mp4", ".wav", ".m4a", ".ogg", ".webm", ".flac"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}
MAX_IMAGES = 4

# ── Rate limiting simples por IP ─────────────────────────────────────────────
from collections import defaultdict
import time

_rate_limit: dict[str, list[float]] = defaultdict(list)
RATE_LIMIT_MAX = 10       # máx requisições
RATE_LIMIT_WINDOW = 60    # por minuto


def _check_rate_limit(ip: str) -> bool:
    """Retorna True se o IP excedeu o limite."""
    agora = time.time()
    _rate_limit[ip] = [t for t in _rate_limit[ip] if agora - t < RATE_LIMIT_WINDOW]
    if len(_rate_limit[ip]) >= RATE_LIMIT_MAX:
        return True
    _rate_limit[ip].append(agora)
    return False

app.config["RESUMOS_FOLDER"].mkdir(exist_ok=True)
app.config["UPLOAD_FOLDER"].mkdir(exist_ok=True)


# ── Supabase ──────────────────────────────────────────────────────────────────
def _supa():
    from supabase import create_client
    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_KEY")
    if not url or not key:
        raise RuntimeError("Variáveis SUPABASE_URL e SUPABASE_KEY devem estar configuradas.")
    return create_client(url, key)


def usuario_logado() -> dict | None:
    """Retorna dados do usuário da sessão ou None."""
    return session.get("usuario")


def salvar_resumo_nuvem(slug: str, titulo: str, data_hora: str,
                        html: str, txt: str, meta: dict) -> bool:
    """Salva o resumo no Supabase Storage e registra na tabela. Retorna True se ok."""
    uid = (usuario_logado() or {}).get("id")
    if not uid:
        return False
    try:
        sb = _supa()
        prefix = f"{uid}/{slug}"
        sb.storage.from_("resumos").upload(
            f"{prefix}.html", html.encode(), {"content-type": "text/html; charset=utf-8", "upsert": "true"})
        sb.storage.from_("resumos").upload(
            f"{prefix}.txt", txt.encode(), {"content-type": "text/plain; charset=utf-8", "upsert": "true"})
        sb.table("resumos").insert({
            "user_id": uid, "slug": slug, "titulo": titulo,
            "data_hora": data_hora, "n_imagens": meta.get("n_imagens", 0),
            "tem_audio": meta.get("tem_audio", False),
            "tipo_url": meta.get("tipo_url", ""),
        }).execute()
        return True
    except Exception as e:
        app.logger.warning(f"Nuvem: {e}")
        return False


def listar_resumos_nuvem() -> list[dict]:
    """Lista resumos do usuário logado."""
    uid = (usuario_logado() or {}).get("id")
    if not uid:
        return []
    try:
        sb = _supa()
        res = sb.table("resumos").select("*") \
            .eq("user_id", uid).order("created_at", desc=True).execute()
        return res.data or []
    except Exception as e:
        app.logger.warning(f"Listar nuvem: {e}")
        return []


def carregar_resumo_nuvem(slug: str) -> str | None:
    """Baixa o HTML de um resumo do Supabase Storage."""
    uid = (usuario_logado() or {}).get("id")
    if not uid:
        return None
    try:
        sb = _supa()
        dados = sb.storage.from_("resumos").download(f"{uid}/{slug}.html")
        return dados.decode("utf-8")
    except Exception:
        return None


def carregar_txt_nuvem(slug: str) -> str | None:
    uid = (usuario_logado() or {}).get("id")
    if not uid:
        return None
    try:
        sb = _supa()
        dados = sb.storage.from_("resumos").download(f"{uid}/{slug}.txt")
        return dados.decode("utf-8")
    except Exception:
        return None


# ── Retry helper para chamadas à API ─────────────────────────────────────────
def _retry(fn, max_tentativas=3, delay=1.0):
    """Executa fn() com até max_tentativas retentativas em caso de erro."""
    for tentativa in range(max_tentativas):
        try:
            return fn()
        except Exception:
            if tentativa == max_tentativas - 1:
                raise
            time.sleep(delay * (tentativa + 1))


# ── Extração de conteúdo de URL ───────────────────────────────────────────────
def _youtube_id(url: str):
    """Extrai o ID do vídeo de uma URL do YouTube (inclui shorts, nocookie, etc.)."""
    patterns = [
        r"(?:v=|youtu\.be/|/embed/|/v/|/shorts/|/live/)([A-Za-z0-9_-]{11})",
        r"youtube-nocookie\.com/embed/([A-Za-z0-9_-]{11})",
        r"^([A-Za-z0-9_-]{11})$",
    ]
    for p in patterns:
        m = re.search(p, url)
        if m:
            return m.group(1)
    return None


def extrair_conteudo_url(url: str) -> tuple[str, str]:
    """
    Retorna (tipo, conteudo):
      tipo: 'youtube' | 'pagina'
      conteudo: texto extraído (truncado)
    """
    import requests

    vid_id = _youtube_id(url)
    if vid_id:
        try:
            from youtube_transcript_api import YouTubeTranscriptApi
            partes = YouTubeTranscriptApi.get_transcript(vid_id, languages=["pt", "pt-BR", "en"])
            texto = " ".join(p["text"] for p in partes)
            return "youtube", texto[:8000]
        except Exception as e:
            return "youtube", f"[Não foi possível obter legenda: {e}]"

    # Página web genérica
    try:
        resp = requests.get(url, timeout=10, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
        # Remove tags HTML
        texto = re.sub(r"<[^>]+>", " ", resp.text)
        texto = re.sub(r"\s+", " ", texto).strip()
        return "pagina", texto[:8000]
    except Exception as e:
        return "pagina", f"[Erro ao acessar o link: {e}]"


# ── Transcrição de áudio via Groq Whisper ─────────────────────────────────────
def transcrever_audio(caminho: str) -> str:
    from groq import Groq
    client = Groq(api_key=os.environ["GROQ_API_KEY"])

    def _call():
        with open(caminho, "rb") as f:
            resp = client.audio.transcriptions.create(
                file=(Path(caminho).name, f),
                model="whisper-large-v3",
                response_format="text",
                language="pt",
            )
        return resp if isinstance(resp, str) else resp.text

    return _retry(_call)


# ── Descrição de imagem via Groq Vision ───────────────────────────────────────
def descrever_imagem(caminho: str) -> str:
    from groq import Groq
    client = Groq(api_key=os.environ["GROQ_API_KEY"])
    with open(caminho, "rb") as f:
        b64 = base64.b64encode(f.read()).decode()
    ext = Path(caminho).suffix.lower().lstrip(".")
    mime = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}"

    def _call():
        resp = client.chat.completions.create(
            model="meta-llama/llama-4-scout-17b-16e-instruct",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:{mime};base64,{b64}"},
                        },
                        {
                            "type": "text",
                            "text": (
                                "Você é um assistente escolar. Descreva de forma detalhada "
                                "o que está nesta imagem de aula: textos, diagramas, fórmulas, "
                                "quadro-negro, slides ou qualquer conteúdo educacional visível."
                            ),
                        },
                    ],
                }
            ],
            max_tokens=500,
        )
        return resp.choices[0].message.content.strip()

    return _retry(_call)


# ── Geração do resumo HTML via Groq ───────────────────────────────────────────
def gerar_resumo(transcricao: str, descricoes_imagens: list[str], titulo: str,
                  conteudo_url: str = "", tipo_url: str = "") -> str:
    from groq import Groq
    client = Groq(api_key=os.environ["GROQ_API_KEY"])

    partes = []
    if transcricao:
        partes.append(f"TRANSCRIÇÃO DO ÁUDIO:\n{transcricao[:4000]}")
    for i, desc in enumerate(descricoes_imagens, 1):
        partes.append(f"IMAGEM {i}:\n{desc[:800]}")
    if conteudo_url:
        label = "TRANSCRIÇÃO DO VÍDEO (YouTube)" if tipo_url == "youtube" else "CONTEÚDO DA PÁGINA"
        partes.append(f"{label}:\n{conteudo_url[:4000]}")

    conteudo = "\n\n".join(partes) if partes else "Nenhum conteúdo enviado."

    def _call():
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Você é um assistente escolar especializado em criar resumos de aulas. "
                        "Organize o conteúdo de forma clara e didática em português brasileiro."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Crie um resumo detalhado e completo da aula '{titulo}' com base neste conteúdo:\n\n"
                        f"{conteudo}\n\n"
                        "Estruture o resumo com:\n"
                        "# Título da aula\n"
                        "## Introdução\n"
                        "## Tópicos Principais (desenvolva cada tópico com detalhes e subtópicos)\n"
                        "## Conceitos-chave (**destaque** os termos importantes em negrito)\n"
                        "## Exemplos e Aplicações\n"
                        "## Conclusão e O que aprender\n"
                        "Use markdown: # para títulos, ## para subtítulos, **negrito** para termos importantes, - para listas.\n"
                        "Seja detalhado e didático."
                    ),
                },
            ],
            max_tokens=2500,
        )
        return resp.choices[0].message.content.strip()

    return _retry(_call)


# ── Converter resumo em HTML ───────────────────────────────────────────────────
def _md_inline(texto: str) -> str:
    """Converte markdown inline (**negrito**, *itálico*) para HTML."""
    texto = re.sub(r"\*\*\*(.+?)\*\*\*", r"<strong><em>\1</em></strong>", texto)
    texto = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", texto)
    texto = re.sub(r"\*(.+?)\*", r"<em>\1</em>", texto)
    texto = re.sub(r"`(.+?)`", r"<code>\1</code>", texto)
    return texto


def resumo_para_html(resumo_texto: str, titulo: str, data_hora: str,
                     n_imagens: int, tem_audio: bool, tipo_url: str = "",
                     slug: str = "") -> str:
    titulo = str(html_escape(titulo))
    linhas = resumo_texto.split("\n")
    blocos = []
    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        stripped = linha.strip()

        if not stripped:
            i += 1
            continue

        # Títulos markdown
        if stripped.startswith("#### "):
            blocos.append(f"<h4>{_md_inline(stripped[5:])}</h4>")
        elif stripped.startswith("### "):
            blocos.append(f"<h3>{_md_inline(stripped[4:])}</h3>")
        elif stripped.startswith("## "):
            blocos.append(f"<h2>{_md_inline(stripped[3:])}</h2>")
        elif stripped.startswith("# "):
            blocos.append(f"<h2>{_md_inline(stripped[2:])}</h2>")
        # Listas com - • *
        elif stripped.startswith(("- ", "• ", "* ")):
            blocos.append(f"<li>{_md_inline(stripped[2:])}</li>")
        # Listas numeradas: "1. " ou "1) "
        elif re.match(r"^\d+[.)]\s", stripped):
            texto = re.sub(r"^\d+[.)]\s+", "", stripped)
            blocos.append(f"<li>{_md_inline(texto)}</li>")
        # Linha toda em maiúsculas = título
        elif stripped.isupper() and len(stripped) < 80:
            blocos.append(f"<h2>{stripped}</h2>")
        else:
            blocos.append(f"<p>{_md_inline(stripped)}</p>")
        i += 1

    corpo = "\n    ".join(blocos)

    slug_placeholder = f"/resumo/{slug}" if slug else "#"
    titulo_safe = titulo.replace("'", "").replace('"', "")

    fontes = []
    if tem_audio:
        fontes.append('<span class="badge audio">🎵 Áudio</span>')
    if n_imagens:
        fontes.append(f'<span class="badge imagem">🖼️ {n_imagens} imagem(ns)</span>')
    if tipo_url == "youtube":
        fontes.append('<span class="badge youtube">▶️ YouTube</span>')
    elif tipo_url == "pagina":
        fontes.append('<span class="badge link">🔗 Link</span>')

    return f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{titulo} — Resumo de Aula</title>
  <link rel="icon" type="image/svg+xml" href="/static/favicon.svg">
  <style>
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{
      font-family: 'Segoe UI', system-ui, sans-serif;
      background: #f0f4f8;
      color: #1a202c;
      padding: 2rem 1rem;
    }}
    .card {{
      max-width: 820px;
      margin: 0 auto;
      background: #fff;
      border-radius: 16px;
      box-shadow: 0 4px 24px rgba(0,0,0,.1);
      overflow: hidden;
    }}
    .header {{
      background: linear-gradient(135deg, #4f46e5, #7c3aed);
      color: #fff;
      padding: 2rem 2.5rem;
    }}
    .header h1 {{ font-size: 1.8rem; margin-bottom: .5rem; }}
    .meta {{ font-size: .85rem; opacity: .85; margin-top: .4rem; }}
    .badges {{ display: flex; gap: .5rem; margin-top: 1rem; flex-wrap: wrap; }}
    .badge {{
      font-size: .78rem;
      padding: .25rem .7rem;
      border-radius: 999px;
      font-weight: 600;
    }}
    .badge.audio {{ background: #fef3c7; color: #92400e; }}
    .badge.imagem {{ background: #dbeafe; color: #1e40af; }}
    .badge.youtube {{ background: #fee2e2; color: #991b1b; }}
    .badge.link {{ background: #d1fae5; color: #065f46; }}
    .body {{ padding: 2rem 2.5rem; line-height: 1.75; }}
    h2 {{
      font-size: 1.2rem;
      color: #4f46e5;
      margin: 1.8rem 0 .6rem;
      border-left: 4px solid #4f46e5;
      padding-left: .75rem;
    }}
    h3 {{
      font-size: 1.05rem;
      color: #374151;
      margin: 1.2rem 0 .4rem;
    }}
    p {{ margin: .5rem 0; color: #374151; }}
    strong {{ color: #1a202c; }}
    h4 {{
      font-size: 1rem;
      color: #6d28d9;
      margin: 1rem 0 .3rem;
    }}
    code {{
      background: #f3f4f6;
      padding: .1rem .35rem;
      border-radius: 4px;
      font-size: .88em;
      color: #7c3aed;
    }}
    li {{
      margin: .4rem 0 .4rem 1.5rem;
      color: #374151;
    }}
    .footer {{
      text-align: center;
      padding: 1.2rem;
      font-size: .78rem;
      color: #9ca3af;
      border-top: 1px solid #f3f4f6;
    }}
    /* Print / PDF */
    @media print {{
      body {{ background: #fff; padding: 0; }}
      .card {{ box-shadow: none; border-radius: 0; }}
      .actions {{ display: none !important; }}
      .header {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
      .ilustracoes {{ break-inside: avoid; }}
      h2 {{ break-after: avoid; }}
    }}
    /* Barra de ações */
    .actions {{
      display: flex;
      gap: .75rem;
      padding: 1.2rem 2.5rem;
      background: #fafafa;
      border-top: 1px solid #f3f4f6;
      flex-wrap: wrap;
    }}
    .btn {{
      display: inline-flex;
      align-items: center;
      gap: .4rem;
      padding: .6rem 1.2rem;
      border-radius: 8px;
      font-size: .88rem;
      font-weight: 600;
      cursor: pointer;
      border: none;
      text-decoration: none;
      transition: opacity .15s, transform .1s;
    }}
    .btn:hover {{ opacity: .88; transform: translateY(-1px); }}
    .btn-primary {{ background: linear-gradient(135deg,#4f46e5,#7c3aed); color:#fff; }}
    .btn-secondary {{ background: #f3f4f6; color: #374151; }}
    .btn-green {{ background: #d1fae5; color: #065f46; }}
    #copy-msg {{ font-size:.8rem; color:#059669; display:none; align-self:center; }}
  </style>
</head>
<body>
  <div class="card">
    <div class="header">
      <h1>📚 {titulo}</h1>
      <div class="meta">Resumo gerado em {data_hora}</div>
      <div class="badges">{' '.join(fontes)}</div>
    </div>
    <div class="body" id="resumo-body">
    {corpo}
    </div>

    <!-- Barra de ações -->
    <div class="actions">
      <button class="btn btn-primary" onclick="baixarHTML()">⬇️ Baixar HTML</button>
      <a class="btn btn-primary" href="{slug_placeholder}/docx" download>📄 Baixar DOCX</a>
      <button class="btn btn-green" onclick="copiarTexto()">📋 Copiar texto</button>
      <button class="btn btn-secondary" onclick="window.print()">🖨️ Salvar PDF</button>
      <a class="btn btn-secondary" href="/">← Voltar</a>
      <span id="copy-msg">Copiado!</span>
    </div>

    <div class="footer">Gerado automaticamente pelo Bot de Resumo de Aulas</div>
  </div>

  <script>
    function baixarHTML() {{
      var html = document.documentElement.outerHTML;
      var blob = new Blob([html], {{type: 'text/html;charset=utf-8'}});
      var a = document.createElement('a');
      a.href = URL.createObjectURL(blob);
      a.download = '{titulo_safe}.html';
      a.click();
      URL.revokeObjectURL(a.href);
    }}
    function copiarTexto() {{
      var el = document.getElementById('resumo-body');
      var texto = el.innerText || el.textContent;
      navigator.clipboard.writeText(texto).then(function() {{
        var msg = document.getElementById('copy-msg');
        msg.style.display = 'inline';
        setTimeout(function() {{ msg.style.display = 'none'; }}, 2000);
      }});
    }}
  </script>
</body>
</html>"""


# ── Helpers DOCX ─────────────────────────────────────────────────────────────
def _docx_inline(paragraph, texto: str):
    """Adiciona runs com negrito/itálico a um parágrafo de DOCX."""
    partes = re.split(r"(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)", texto)
    for parte in partes:
        if parte.startswith("***") and parte.endswith("***"):
            r = paragraph.add_run(parte[3:-3])
            r.bold = True
            r.italic = True
        elif parte.startswith("**") and parte.endswith("**"):
            r = paragraph.add_run(parte[2:-2])
            r.bold = True
        elif parte.startswith("*") and parte.endswith("*"):
            r = paragraph.add_run(parte[1:-1])
            r.italic = True
        else:
            paragraph.add_run(parte)


# ── Auth ──────────────────────────────────────────────────────────────────────
@app.route("/auth/google")
def auth_google():
    try:
        sb = _supa()
        r = sb.auth.sign_in_with_oauth({
            "provider": "google",
            "options": {
                "redirect_to": request.host_url.rstrip("/") + "/auth/callback"
            }
        })
        return redirect(r.url)
    except Exception as e:
        flash(f"Erro ao iniciar login com Google: {e}", "erro")
        return redirect(url_for("entrar"))


@app.route("/auth/callback")
def auth_callback():
    """Página que extrai o token do fragment (#) via JS e envia ao servidor."""
    return render_template("auth_callback.html")


@app.route("/auth/session", methods=["POST"])
@csrf.exempt
def auth_session():
    """Recebe access_token + refresh_token do JS e cria a sessão Flask."""
    data = request.get_json(force=True) or {}
    access_token  = data.get("access_token", "")
    refresh_token = data.get("refresh_token", "")
    if not access_token:
        return jsonify({"ok": False}), 400
    try:
        sb = _supa()
        r = sb.auth.set_session(access_token, refresh_token)
        user = r.user
        session["usuario"] = {
            "id":            user.id,
            "email":         user.email,
            "access_token":  access_token,
            "refresh_token": refresh_token,
        }
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "erro": str(e)}), 400


@app.route("/entrar", methods=["GET", "POST"])
def entrar():
    if usuario_logado():
        return redirect(url_for("index"))
    erro = None
    if request.method == "POST":
        email = request.form.get("email", "").strip()
        senha = request.form.get("senha", "")
        acao  = request.form.get("acao", "login")  # "login" ou "cadastro"
        try:
            sb = _supa()
            if acao == "cadastro":
                r = sb.auth.sign_up({"email": email, "password": senha})
                if r.user:
                    session["usuario"] = {"id": r.user.id, "email": r.user.email,
                                          "access_token": r.session.access_token,
                                          "refresh_token": r.session.refresh_token}
                    return redirect(url_for("index"))
                erro = "Não foi possível criar a conta."
            else:
                r = sb.auth.sign_in_with_password({"email": email, "password": senha})
                if r.user:
                    session["usuario"] = {"id": r.user.id, "email": r.user.email,
                                          "access_token": r.session.access_token,
                                          "refresh_token": r.session.refresh_token}
                    return redirect(url_for("index"))
                erro = "E-mail ou senha incorretos."
        except Exception as e:
            msg = str(e)
            if "Invalid login" in msg or "invalid_credentials" in msg:
                erro = "E-mail ou senha incorretos."
            elif "already registered" in msg:
                erro = "E-mail já cadastrado. Faça login."
            elif "password" in msg.lower():
                erro = "A senha deve ter pelo menos 6 caracteres."
            elif "connection" in msg.lower() or "timeout" in msg.lower():
                erro = "Serviço indisponível. Tente novamente em alguns instantes."
            else:
                erro = "Erro ao autenticar. Tente novamente."
    return render_template("auth.html", erro=erro)


@app.route("/sair")
def sair():
    session.clear()
    return redirect(url_for("index"))


# ── Rotas principais ──────────────────────────────────────────────────────────
@app.route("/")
def index():
    usuario = usuario_logado()
    resumos = listar_resumos_nuvem() if usuario else []
    return render_template("index.html", resumos=resumos, usuario=usuario)


@app.route("/upload", methods=["POST"])
def upload():
    if _check_rate_limit(request.remote_addr or "unknown"):
        return jsonify({"erro": "Muitas requisições. Aguarde um momento."}), 429

    titulo = request.form.get("titulo", "Aula sem título").strip() or "Aula sem título"
    link   = request.form.get("link", "").strip()
    arquivos = request.files.getlist("arquivos")

    sem_arquivos = not arquivos or all(f.filename == "" for f in arquivos)
    if sem_arquivos and not link:
        return jsonify({"erro": "Envie pelo menos um arquivo ou um link."}), 400

    ALLOWED_EXTS = AUDIO_EXTS | IMAGE_EXTS
    audios, imagens = [], []
    for arq in arquivos:
        if not arq.filename:
            continue
        nome = secure_filename(arq.filename)
        ext = Path(nome).suffix.lower()
        if ext not in ALLOWED_EXTS:
            return jsonify({"erro": f"Tipo de arquivo não suportado: {ext}"}), 400
        destino = app.config["UPLOAD_FOLDER"] / nome
        arq.save(str(destino))
        if ext in AUDIO_EXTS:
            audios.append(str(destino))
        elif ext in IMAGE_EXTS:
            imagens.append(str(destino))

    transcricao = ""
    if audios:
        try:
            transcricao = transcrever_audio(audios[0])
        except Exception as e:
            transcricao = f"[Erro na transcrição: {e}]"

    imagens_ignoradas = max(0, len(imagens) - MAX_IMAGES)
    descricoes = []
    for img in imagens[:MAX_IMAGES]:
        try:
            descricoes.append(descrever_imagem(img))
        except Exception as e:
            descricoes.append(f"[Erro ao descrever imagem: {e}]")

    conteudo_url, tipo_url = "", ""
    if link:
        try:
            tipo_url, conteudo_url = extrair_conteudo_url(link)
        except Exception as e:
            conteudo_url = f"[Erro ao processar link: {e}]"
            tipo_url = "pagina"

    try:
        resumo_texto = gerar_resumo(transcricao, descricoes, titulo, conteudo_url, tipo_url)
    except Exception as e:
        resumo_texto = f"Erro ao gerar resumo: {e}"

    agora    = datetime.now()
    slug     = agora.strftime("%Y%m%d_%H%M%S")
    data_hora = agora.strftime("%d/%m/%Y às %H:%M")
    meta     = {"titulo": titulo, "slug": slug, "data_hora": data_hora,
                "n_imagens": len(imagens), "tem_audio": bool(audios), "tipo_url": tipo_url,
                "imagens_ignoradas": imagens_ignoradas}
    if imagens_ignoradas:
        flash(f"Apenas {MAX_IMAGES} imagens foram processadas. {imagens_ignoradas} imagem(ns) foram ignoradas.", "aviso")

    html_content = resumo_para_html(resumo_texto, titulo, data_hora,
                                    len(imagens), bool(audios), tipo_url, slug)

    # Salva localmente (sempre) + nuvem (se logado)
    pasta = app.config["RESUMOS_FOLDER"]
    (pasta / f"resumo_{slug}.html").write_text(html_content, encoding="utf-8")
    (pasta / f"resumo_{slug}.txt").write_text(resumo_texto, encoding="utf-8")
    (pasta / f"resumo_{slug}.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    salvar_resumo_nuvem(slug, titulo, data_hora, html_content, resumo_texto, meta)

    # Limpa arquivos de upload temporários
    for caminho in audios + imagens:
        try:
            os.remove(caminho)
        except OSError:
            pass

    return redirect(url_for("ver_resumo", slug=slug))


@app.route("/resumo/<slug>")
def ver_resumo(slug: str):
    # Tenta local primeiro, depois nuvem
    pasta = app.config["RESUMOS_FOLDER"]
    html_path = pasta / f"resumo_{slug}.html"
    if html_path.exists():
        return html_path.read_text(encoding="utf-8")
    html = carregar_resumo_nuvem(slug)
    if html:
        return html
    return "Resumo não encontrado.", 404


@app.route("/resumo/<slug>/deletar", methods=["POST"])
def deletar_resumo(slug: str):
    uid = (usuario_logado() or {}).get("id")
    if not uid:
        return jsonify({"erro": "Não autorizado."}), 401
    try:
        sb = _supa()
        # Remove do Storage
        prefix = f"{uid}/{slug}"
        sb.storage.from_("resumos").remove([f"{prefix}.html", f"{prefix}.txt"])
        # Remove do banco
        sb.table("resumos").delete().eq("user_id", uid).eq("slug", slug).execute()
    except Exception as e:
        app.logger.warning(f"Erro ao deletar resumo: {e}")
    # Remove arquivos locais
    pasta = app.config["RESUMOS_FOLDER"]
    for ext in (".html", ".txt", ".json"):
        try:
            (pasta / f"resumo_{slug}{ext}").unlink(missing_ok=True)
        except OSError:
            pass
    flash("Resumo deletado com sucesso.", "sucesso")
    return redirect(url_for("index"))


@app.route("/resumo/<slug>/docx")
def baixar_docx(slug: str):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from io import BytesIO
    from flask import send_file

    pasta = app.config["RESUMOS_FOLDER"]
    txt_path  = pasta / f"resumo_{slug}.txt"
    json_path = pasta / f"resumo_{slug}.json"

    resumo_texto = txt_path.read_text(encoding="utf-8") if txt_path.exists() \
                   else carregar_txt_nuvem(slug)
    if not resumo_texto:
        return "Resumo não encontrado.", 404

    meta      = json.loads(json_path.read_text(encoding="utf-8")) if json_path.exists() else {}
    titulo    = meta.get("titulo", "Resumo de Aula")
    data_hora = meta.get("data_hora", "")

    doc = Document()
    t = doc.add_heading(titulo, level=0)
    t.runs[0].font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
    if data_hora:
        p = doc.add_paragraph(f"Gerado em {data_hora}")
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    doc.add_paragraph("")

    for linha in resumo_texto.split("\n"):
        s = linha.strip()
        if not s:
            doc.add_paragraph(""); continue
        if s.startswith("#### "):   doc.add_heading(s[5:], level=4)
        elif s.startswith("### "): doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):  doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):   doc.add_heading(s[2:], level=1)
        elif s.startswith(("- ","• ","* ")):
            _docx_inline(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\d+[.)]\s", s):
            _docx_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+","",s))
        else:
            _docx_inline(doc.add_paragraph(), s)

    buf = BytesIO()
    doc.save(buf); buf.seek(0)
    nome = re.sub(r"[^\w\s-]", "", titulo)[:50].strip() + ".docx"
    return send_file(buf, as_attachment=True, download_name=nome,
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.route("/resumos/<path:filename>")
def static_resumos(filename):
    return send_from_directory(app.config["RESUMOS_FOLDER"], filename)


if __name__ == "__main__":
    app.run(debug=True, port=5000)
